"""Run with CODEX_PRIMARY_RUNTIME_PYTHON; generates the two website Word templates."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]

def rtl(p):
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    pPr=p._p.get_or_add_pPr();b=OxmlElement('w:bidi');pPr.append(b)
    return p

def paragraph(d,label,value='',style=None,keep=False):
    p=rtl(d.add_paragraph(style=style));p.paragraph_format.keep_with_next=keep
    p.paragraph_format.space_after=Pt(3)
    for text,bold,color in [(label,True,'000000'),(value,False,'666666')]:
        r=p.add_run(text);r.bold=bold;r.font.name='Arial';r.font.color.rgb=RGBColor.from_string(color)
        props=r._element.get_or_add_rPr();fonts=props.rFonts
        if fonts is not None:fonts.set(qn('w:cs'),'Arial')
        lang=OxmlElement('w:rtl');props.append(lang)
    return p

def build(count,output,filled=False):
    d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(1.6);sec.left_margin=sec.right_margin=Cm(1.8)
    for key in ['Normal','Title','Heading 1']:
        st=d.styles[key];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(12 if key=='Normal' else 20 if key=='Title' else 15)
        st.paragraph_format.line_spacing=1.05
    for node in d.styles.element.xpath('.//w:pBdr'):
        node.getparent().remove(node)
    d.core_properties.title='قالب إعداد أسئلة الاختبار';d.core_properties.author='ورقة'
    paragraph(d,'قالب إعداد أسئلة الاختبار',style='Title')
    paragraph(d,'',f'{count} سؤال اختيار من متعدد و{count} سؤال صح وخطأ')
    paragraph(d,'','استبدل النصوص الإرشادية بأسئلتك، وأبق عناوين الأسئلة وحروف الخيارات. احفظ الملف من Word بصيغة PDF، ثم ارفعه عبر استيراد أسئلة PDF في ورقة. مفتاح الإجابة اختياري ويمكن ضبطه في الموقع.')
    paragraph(d,'اختيار من متعدد',style='Heading 1',keep=True)
    for i in range(1,count+1):
        if i>1 and (i-1)%5==0:
            d.add_page_break();paragraph(d,'اختيار من متعدد',style='Heading 1',keep=True)
        paragraph(d,f'اختياري {i:02d}: ','أي الكواكب أقرب إلى الشمس؟' if filled else 'اكتب نص السؤال هنا',keep=True)
        for letter,val in zip('أبجد',['عطارد','الأرض','المريخ','زحل']):
            paragraph(d,f'{letter}: ',val if filled else f'اكتب الخيار {letter} هنا',keep=True)
        p=paragraph(d,'الإجابة الصحيحة: ',('أ' if filled else '...')+'    الدرجة: 1');p.paragraph_format.space_after=Pt(8)
    d.add_page_break();paragraph(d,'صح وخطأ',style='Heading 1',keep=True)
    for i in range(1,count+1):
        paragraph(d,f'صح وخطأ {i:02d}: ','تدور الأرض حول الشمس.' if filled else 'اكتب العبارة هنا',keep=True)
        p=paragraph(d,'الإجابة الصحيحة: ',('صح' if filled else '...')+'    الدرجة: 1');p.paragraph_format.space_after=Pt(8)
    d.save(output)

if __name__=='__main__':
    for n in [10,15]:build(n,ROOT/f'public/templates/questions-{n}-{n}.docx')
